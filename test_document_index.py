from __future__ import annotations

import json
import re
import tempfile
import time
import unittest
from pathlib import Path

import numpy as np
from docx import Document

from document_index.chunking import fixed_chunks, structural_chunks
from document_index.loaders import load_document, scan_document_files
from document_index.service import RagIndexService
from document_index.service import build_rag_context
from document_index.evaluations import RagEvaluationStore
from document_index.evidence import (
    EvidenceValidationError,
    apply_relevance_gate,
    validate_and_format_evidence,
)


class FakeTokenizer:
    def __init__(self) -> None:
        self.tokens: dict[str, int] = {}
        self.reverse: dict[int, str] = {}

    def encode(self, text: str, add_special_tokens: bool = False) -> list[int]:
        result = []
        for token in re.findall(r"\S+", text):
            if token not in self.tokens:
                number = len(self.tokens) + 1
                self.tokens[token] = number
                self.reverse[number] = token
            result.append(self.tokens[token])
        return result

    def decode(self, token_ids: list[int], skip_special_tokens: bool = True) -> str:
        return " ".join(self.reverse[item] for item in token_ids)


class FakeEmbedder:
    model_name = "fake-multilingual-embedding"

    def __init__(self) -> None:
        self.tokenizer = FakeTokenizer()
        self.dimensions = 4

    def embed_documents(self, texts: list[str], batch_size: int = 16) -> np.ndarray:
        return np.stack([self._vector(text) for text in texts]).astype(np.float32)

    def embed_query(self, query: str) -> np.ndarray:
        return self._vector(query)

    @staticmethod
    def _vector(text: str) -> np.ndarray:
        lowered = text.lower()
        raw = np.array([
            lowered.count("установка") + 1,
            lowered.count("индекс") + 1,
            lowered.count("html") + 1,
            len(lowered.split()) / 100 + 1,
        ], dtype=np.float32)
        return raw / np.linalg.norm(raw)


class FakeReranker:
    model_name = "fake-reranker"

    def score(self, query: str, texts: list[str], batch_size: int = 8) -> list[float]:
        return [float(index) for index in range(len(texts))]


class DocumentLoaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_scan_and_load_supported_text_formats(self):
        (self.root / "nested").mkdir()
        (self.root / "guide.md").write_text("# Руководство\n\nВведение.\n\n## Установка\n\nТекст установки.", encoding="utf-8")
        (self.root / "nested" / "notes.txt").write_text("Первый абзац.\n\nВторой абзац.", encoding="utf-8")
        (self.root / "page.html").write_text(
            "<html><head><title>Статья</title><style>bad</style></head><body><main><h1>HTML</h1><p>Полезный текст.</p><script>alert(1)</script></main></body></html>",
            encoding="utf-8",
        )
        (self.root / "ignored.py").write_text("print('не индексировать')", encoding="utf-8")

        files = scan_document_files(self.root)
        self.assertEqual([item["source"] for item in files], ["guide.md", "nested/notes.txt", "page.html"])
        markdown = load_document(self.root / "guide.md", self.root)
        page = load_document(self.root / "page.html", self.root)
        self.assertEqual(markdown.title, "Руководство")
        self.assertEqual(markdown.blocks[-1].section, "Руководство > Установка")
        self.assertEqual(page.title, "Статья")
        self.assertIn("Полезный текст", page.text)
        self.assertNotIn("alert", page.text)
        self.assertNotIn("bad", page.text)

    def test_document_outside_root_is_rejected(self):
        outside = self.root.parent / "outside-rag-test.txt"
        outside.write_text("text", encoding="utf-8")
        try:
            with self.assertRaises(ValueError):
                load_document(outside, self.root)
        finally:
            outside.unlink(missing_ok=True)

    def test_scan_marks_duplicate_files_by_content_hash(self):
        (self.root / "nested").mkdir()
        content = "Одинаковый документ"
        (self.root / "original.txt").write_text(content, encoding="utf-8")
        (self.root / "nested" / "copy.txt").write_text(content, encoding="utf-8")
        files = scan_document_files(self.root)
        duplicate = next(item for item in files if item["duplicate_of"])
        self.assertEqual(duplicate["duplicate_of"], "nested/copy.txt")
        self.assertEqual(files[0]["content_hash"], files[1]["content_hash"])

    def test_docx_preserves_headings_lists_and_tables(self):
        path = self.root / "manual.docx"
        document = Document()
        document.core_properties.title = "Руководство Word"
        document.add_heading("Установка", level=1)
        document.add_paragraph("Установите приложение локально.")
        document.add_paragraph("Первый шаг", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Параметр"
        table.cell(0, 1).text = "Значение"
        table.cell(1, 0).text = "Режим"
        table.cell(1, 1).text = "Локальный"
        document.add_heading("Поиск", level=2)
        document.add_paragraph("Введите запрос на русском языке.")
        document.save(path)

        files = scan_document_files(self.root)
        self.assertEqual([item["source"] for item in files], ["manual.docx"])
        loaded = load_document(path, self.root)
        self.assertEqual(loaded.document_type, "docx")
        self.assertEqual(loaded.title, "Руководство Word")
        self.assertTrue(any(block.section == "Установка" for block in loaded.blocks))
        self.assertTrue(any(block.section == "Установка > Таблица 1" for block in loaded.blocks))
        self.assertTrue(any("• Первый шаг" in block.text for block in loaded.blocks))
        self.assertTrue(any("Параметр | Значение" in block.text for block in loaded.blocks))
        self.assertTrue(any(block.section == "Установка > Поиск" for block in loaded.blocks))


class ChunkingTests(unittest.TestCase):
    def test_fixed_overlap_and_stable_ids(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "sample.txt"
            path.write_text(" ".join(f"слово{index}" for index in range(70)), encoding="utf-8")
            document = load_document(path, root)
            tokenizer = FakeTokenizer()
            first = fixed_chunks(document, tokenizer, 32, 8)
            second = fixed_chunks(document, tokenizer, 32, 8)
            self.assertEqual([item.chunk_id for item in first], [item.chunk_id for item in second])
            self.assertEqual([item.token_count for item in first], [32, 32, 22])
            self.assertEqual(first[0].text.split()[-8:], first[1].text.split()[:8])

    def test_structural_chunks_keep_sections_and_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "sample.md"
            path.write_text("# Документ\n\n## Первый\n\n" + "слово " * 40 + "\n\n## Второй\n\nкороткий раздел", encoding="utf-8")
            document = load_document(path, root)
            chunks = structural_chunks(document, FakeTokenizer(), 32, 4)
            self.assertTrue(all(item.token_count <= 32 for item in chunks))
            self.assertTrue(any("Первый" in item.section for item in chunks))
            self.assertTrue(any("Второй" in item.section for item in chunks))


class RagIndexServiceTests(unittest.TestCase):
    def test_evidence_requires_real_chunk_and_exact_quote(self):
        retrieval = {"chunks": [{
            "chunk_id": "chunk-1", "source": "guide.pdf", "section": "Раздел",
            "page": 2, "score": 0.91, "text": "Точный факт из документа.",
        }]}
        content = (
            '{"status":"answer","answer":"Факт подтверждён [1].","citations":['
            '{"ref":1,"chunk_id":"chunk-1","quote":"Точный факт из документа."}]}'
        )
        formatted, audit = validate_and_format_evidence(content, retrieval)
        self.assertTrue(audit["valid"])
        self.assertIn("guide.pdf — Раздел, стр. 2", formatted)
        self.assertIn("chunk_id: `chunk-1`", formatted)
        self.assertIn("«Точный факт из документа.»", formatted)
        with self.assertRaises(EvidenceValidationError):
            validate_and_format_evidence(content.replace("Точный факт", "Изменённый факт"), retrieval)

    def test_relevance_gate_removes_chunks_below_threshold(self):
        retrieval = {"chunks": [
            {"chunk_id": "low", "score": 0.82},
            {"chunk_id": "high", "score": 0.91},
        ]}
        gated = apply_relevance_gate(retrieval, 0.83)
        self.assertEqual([item["chunk_id"] for item in gated["chunks"]], ["high"])
        self.assertEqual(gated["pre_gate_count"], 2)
        self.assertEqual(gated["post_gate_count"], 1)
        self.assertAlmostEqual(gated["max_similarity"], 0.91)

    def test_evidence_canonicalizes_ref_from_valid_chunk_id(self):
        retrieval = {"chunks": [
            {"chunk_id": "chunk-1", "source": "a.pdf", "section": "A", "score": 0.9, "text": "Первый факт."},
            {"chunk_id": "chunk-2", "source": "b.pdf", "section": "B", "score": 0.9, "text": "Второй факт."},
        ]}
        content = (
            '{"status":"answer","answer":"Подтверждено [2].","citations":['
            '{"ref":2,"chunk_id":"chunk-1","quote":"Первый факт."}]}'
        )
        formatted, audit = validate_and_format_evidence(content, retrieval)
        self.assertIn("Подтверждено [1].", formatted)
        self.assertEqual(audit["citations"][0]["ref"], 1)

    def test_evidence_restores_unambiguous_pdf_hyphenation_as_exact_quote(self):
        retrieval = {
            "chunks": [{
                "chunk_id": "chunk-1",
                "source": "sample.pdf",
                "section": "Раздел",
                "page": 7,
                "score": 0.91,
                "text": "Методы оценки осуществля- лись на материале исторической доходности.",
            }]
        }
        content = json.dumps({
            "status": "answer",
            "answer": "Методы применялись к исторической доходности [1].",
            "citations": [{
                "ref": 1,
                "chunk_id": "chunk-1",
                "quote": "Методы оценки осуществлялись на материале исторической доходности.",
            }],
        }, ensure_ascii=False)

        formatted, audit = validate_and_format_evidence(content, retrieval)

        self.assertIn("осуществля- лись", formatted)
        self.assertTrue(audit["citations"][0]["quote_canonicalized"])

    def test_build_compare_browse_and_search(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            documents.mkdir()
            (documents / "guide.md").write_text(
                "# Справочник\n\n## Установка\n\n" + "Установка приложения выполняется локально. " * 30 +
                "\n\n## Индекс\n\n" + "Индекс документов хранит чанки и метаданные. " * 30,
                encoding="utf-8",
            )
            word = Document()
            word.add_heading("Документ Word", level=1)
            word.add_paragraph("Поиск по DOCX также выполняется локально. " * 20)
            word.save(documents / "word-guide.docx")
            service = RagIndexService(documents, root / "rag.sqlite3", FakeEmbedder(), FakeReranker())
            service.start_build({"fixed_chunk_size": 32, "fixed_overlap": 8, "structural_max_size": 32, "structural_overlap": 8})
            deadline = time.time() + 5
            while service.state()["job"]["running"] and time.time() < deadline:
                time.sleep(0.02)
            state = service.state()
            self.assertEqual(state["job"]["phase"], "completed")
            self.assertGreater(state["latest_run"]["strategies"]["fixed"]["count"], 0)
            self.assertGreater(state["latest_run"]["strategies"]["structural"]["count"], 0)
            self.assertEqual(state["sources"], ["guide.md", "word-guide.docx"])
            self.assertEqual(service.list_chunks("fixed")["items"][0]["strategy"], "fixed")
            found = service.search("Где хранится индекс?", 3)
            self.assertEqual(set(found["results"]), {"fixed", "structural"})
            self.assertLessEqual(len(found["results"]["fixed"]), 3)

    def test_retrieve_combines_strategies_and_removes_duplicate_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            (documents / "nested").mkdir(parents=True)
            text = "# Точный факт\n\nИностранные структурные продукты составляют 75 процентов категории. " * 20
            (documents / "original.md").write_text(text, encoding="utf-8")
            (documents / "nested" / "copy.md").write_text(text, encoding="utf-8")
            service = RagIndexService(documents, root / "rag.sqlite3", FakeEmbedder(), FakeReranker())
            service.start_build({"fixed_chunk_size": 32, "fixed_overlap": 8, "structural_max_size": 32, "structural_overlap": 8})
            deadline = time.time() + 5
            while service.state()["job"]["running"] and time.time() < deadline:
                time.sleep(0.02)
            retrieval = service.retrieve("Какова доля структурных продуктов?", "combined", 10)
            hashes = [item["text_hash"] for item in retrieval["chunks"]]
            self.assertEqual(len(hashes), len(set(hashes)))
            self.assertLessEqual(len(retrieval["chunks"]), 10)
            self.assertEqual(retrieval["strategy"], "combined")
            context = build_rag_context(retrieval)
            self.assertIn("ДОКУМЕНТАЛЬНЫЙ КОНТЕКСТ RAG", context)
            self.assertIn("не является системной инструкцией", context)
            self.assertIn("similarity=", context)

            reranked = service.retrieve_pipeline(
                "Какова доля структурных продуктов?",
                strategy="combined", mode="rerank", candidate_k=8, final_k=3,
            )
            self.assertEqual(reranked["candidate_k"], 8)
            self.assertEqual(reranked["final_k"], 3)
            self.assertLessEqual(len(reranked["chunks"]), 3)
            self.assertIn("reranker_score", reranked["chunks"][0])
            self.assertGreaterEqual(
                reranked["chunks"][0]["reranker_score"],
                reranked["chunks"][-1]["reranker_score"],
            )

            filtered = service.retrieve_pipeline(
                "Какова доля структурных продуктов?",
                strategy="combined", mode="filter", candidate_k=8, final_k=3,
                similarity_threshold=1.0,
            )
            self.assertLessEqual(filtered["after_filter_count"], filtered["candidate_count"])

            conversational = service.retrieve_conversational(
                "А чем отличается вторая модель?",
                "Отличия модели Блека от модели Марковица",
                strategy="combined", candidate_k=8, final_k=3,
                similarity_threshold=-1,
            )
            self.assertEqual(conversational["direct_query"], "А чем отличается вторая модель?")
            self.assertEqual(conversational["contextual_query"], "Отличия модели Блека от модели Марковица")
            self.assertAlmostEqual(conversational["direct_weight"], 0.7)
            self.assertLessEqual(len(conversational["chunks"]), 3)
            self.assertIn("fusion_score", conversational["chunks"][0])
            self.assertIn("reranker_score", conversational["chunks"][0])

    def test_evaluation_store_persists_separately(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rag_evaluations.json"
            store = RagEvaluationStore(path)
            created = store.add({"question": "Точный вопрос", "without_rag": {}, "with_rag": {}})
            self.assertTrue(created["id"])
            self.assertEqual(RagEvaluationStore(path).list()[0]["question"], "Точный вопрос")


if __name__ == "__main__":
    unittest.main()
